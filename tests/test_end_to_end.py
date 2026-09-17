"""End-to-end integration tests (v3).

Run: python tests/test_end_to_end.py
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from config import settings                                     # noqa: E402
from engine import input_parser, retrieval, rules, suggest       # noqa: E402
from indexer.extractor import extract                            # noqa: E402
from output import excel_writer                                  # noqa: E402

FAILURES = []


def check(name, condition, detail=""):
    if condition:
        print(f"  PASS  {name}")
    else:
        print(f"  FAIL  {name} {detail}")
        FAILURES.append(name)


ARTICLE_HTML = """
<html><head>
<title>Immunotherapy Advances in Kidney Cancer | Binaytara | The Cancer News</title>
<meta name="description" content="How checkpoint inhibitors changed outcomes for advanced kidney cancer.">
<link rel="canonical" href="https://binaytara.org/cancernews/article/test-fixture">
</head><body>
<nav><a href="https://binaytara.org/about">About Us</a>
     <a href="https://binaytara.org/research-grants">Funding</a></nav>
<article>
<h1>Immunotherapy Advances in Kidney Cancer</h1>
<p>Kidney cancer affects hundreds of thousands of people every single year across the world and the numbers continue to climb steadily.</p>
<p><strong>Key Takeaways</strong></p>
<p>Checkpoint inhibitors have changed outcomes for patients with advanced kidney cancer across many different treatment settings around the world today.</p>
<p><strong>Background</strong></p>
<p>Renal cell carcinoma accounts for the large majority of kidney cancer diagnoses worldwide, and its incidence has continued to rise steadily across many countries over recent decades of cancer surveillance work.</p>
<p>Immune checkpoint inhibitor combinations are now considered standard first line therapy for many patients living with advanced renal cell carcinoma, and randomised trials show improved overall survival versus older targeted agents.</p>
<p>&ldquo;Early detection is a major gap in the field,&rdquo; said Dr. Yousef Zakharia, noting that far too many patients still present with metastatic disease at the time of their initial diagnosis today.</p>
<p>This paragraph already links to the <a href="https://binaytara.org/research-grants">Binaytara Research Grants</a> page, and it is long enough to be eligible for further suggestions under the rules.</p>
<p><strong>Works discussed</strong></p>
<p>Motzer RJ and colleagues. Kidney cancer NCCN guidelines version three, Journal of the National Comprehensive Cancer Network, volume twenty, pages seventy-one to ninety.</p>
</article>
<footer><a href="https://binaytara.org/contact">Contact</a></footer>
</body></html>
"""

FIXTURE_URL = "https://binaytara.org/cancernews/article/test-fixture"


def make_docx() -> bytes:
    import docx
    d = docx.Document()
    d.add_heading("Immunotherapy Advances in Kidney Cancer", 1)
    d.add_heading("Key Takeaways", 2)
    d.add_paragraph("Checkpoint inhibitors have changed outcomes for patients with "
                    "advanced kidney cancer across many different treatment settings "
                    "around the world in recent years of clinical practice.")
    d.add_heading("Background", 2)
    d.add_paragraph("Renal cell carcinoma accounts for the large majority of kidney "
                    "cancer diagnoses worldwide, and its incidence has continued to "
                    "rise steadily across many countries over recent decades of "
                    "cancer surveillance and registry reporting work.")
    d.add_paragraph("Immune checkpoint inhibitor combinations are now considered "
                    "standard first line therapy for many patients living with "
                    "advanced renal cell carcinoma, and randomised trials show "
                    "improved overall survival versus older targeted agents.")
    d.add_paragraph('"Early detection is a major gap in the field," said Dr. Yousef '
                    "Zakharia, noting that far too many patients still present with "
                    "metastatic disease at the time of their initial diagnosis.")
    d.add_heading("Works discussed", 2)
    d.add_paragraph("Motzer RJ and colleagues. Kidney cancer NCCN guidelines version "
                    "three, JNCCN, volume twenty, pages seventy-one to ninety.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def article_from_html() -> dict:
    rec = extract(ARTICLE_HTML, FIXTURE_URL)
    blocks = rec.pop("_blocks")
    rec.pop("_body_text", None)
    from indexer.chunker import chunk_page
    rec["chunks"] = chunk_page(rec["url"], blocks)
    rec["blocks"] = blocks
    rec["is_draft"] = False
    return rec


def test_artifacts():
    print("\nArtifact load")
    missing = [n for n in ("faiss.index", "pages.json", "paragraphs.json",
                           "manifest.json") if not (settings.DATA / n).exists()]
    check("all index artifacts present", not missing, missing)
    if missing:
        print("  (skipping the rest; run python -m indexer.build_index first)")
        return None
    store = retrieval.load()
    check("FAISS vectors match chunk count",
          store.index.ntotal == len(store.chunks),
          f"{store.index.ntotal} vs {len(store.chunks)}")
    check("manifest records the embedding contract",
          store.manifest.get("model_name") == settings.MODEL_NAME
          and int(store.manifest.get("embedding_dimension", 0)) == settings.EMBED_DIM)
    check("pages present", len(store.pages) > 0, len(store.pages))
    check("every indexed page URL is suggestible",
          all(rules.url_ok(u) for u in list(store.pages)[:200]))
    check("no subdomain leaked into the index",
          not any(".binaytara.org" in u or "binayfoundation" in u for u in store.pages))
    # v3: body texts loaded.
    check("body texts loaded", len(store.body_texts) > 0 or
          not (settings.DATA / "body_texts.json").exists(),
          f"body_texts={len(store.body_texts)}")
    return store


def test_html_pipeline(store):
    print("\nHTML article pipeline")
    art = article_from_html()
    check("canonical URL used", art["url"] == FIXTURE_URL, art["url"])
    check("brand stripped from title",
          "Binaytara" not in art["title_clean"], art["title_clean"])
    check("body link captured",
          "https://binaytara.org/research-grants" in art["body_internal_links"])
    check("nav link NOT counted as a body link",
          "https://binaytara.org/about" not in art["body_internal_links"])
    check("footer link NOT counted as a body link",
          "https://binaytara.org/contact" not in art["body_internal_links"])
    check("chunks produced", len(art["chunks"]) > 0, len(art["chunks"]))

    reasons = {b["skip_reason"] for b in art["blocks"]}
    check("Key Takeaways detected via bold pseudo-heading", "KEY_TAKEAWAYS" in reasons,
          reasons)
    check("references detected via bold pseudo-heading", "REFERENCES" in reasons, reasons)
    check("physician quote detected", "QUOTE" in reasons, reasons)

    res = suggest.analyse(art, store)
    check("analyse returns both tables", "give" in res and "receive" in res)
    check("R7 never suggests an already-linked target",
          all(r["target_url"] != "https://binaytara.org/research-grants"
              for r in res["give"]))
    check("no self-links", all(r["target_url"] != art["url"] for r in res["give"]))
    check("no duplicate targets",
          len({r["target_url"] for r in res["give"]}) == len(res["give"]))
    per_block = {}
    for r in res["give"]:
        per_block[r["block_index"]] = per_block.get(r["block_index"], 0) + 1
    check("max two suggestions per paragraph",
          all(v <= settings.MAX_PER_PARAGRAPH for v in per_block.values()), per_block)
    check("every anchor is 2 to 5 words",
          all(rules.anchor_word_count_ok(r["anchor"]) for r in res["give"]))
    check("every target URL is clean",
          all("?" not in r["target_url"] and "#" not in r["target_url"]
              for r in res["give"]))
    check("every emitted URL is absolute https",
          all(r["target_url"].startswith("https://binaytara.org/")
              for r in res["give"]))
    check("suggestions never land in a skipped block",
          all(any(b["index"] == r["block_index"] and b["eligible"]
                  for b in art["blocks"]) for r in res["give"]))
    # v3: check LLM metadata is present.
    check("analyse result includes LLM status", "llm" in res, list(res.keys()))
    return res


def test_docx_pipeline(store):
    print("\nDOCX draft pipeline")
    art = input_parser.from_docx(make_docx(), "draft.docx")
    reasons = {b.skip_reason for b in art["blocks"]}
    check("Key Takeaways skipped in draft", "KEY_TAKEAWAYS" in reasons, reasons)
    check("references skipped in draft", "REFERENCES" in reasons, reasons)
    check("quote skipped in draft", "QUOTE" in reasons, reasons)
    check("draft flagged as draft", art["is_draft"])
    res = suggest.analyse(art, store)
    leaked = [r for r in res["give"] + res["receive"]
              if rules.DRAFT_SCHEME in r["modified_sentence"]]
    check("draft:// placeholder never reaches the writer", not leaked,
          leaked[0]["modified_sentence"][-60:] if leaked else "")
    check("receive rows warn the article is unpublished",
          all("not yet published" in r["notes"] for r in res["receive"])
          if res["receive"] else True)
    return res


def test_workbook(results):
    print("\nExcel workbook")
    data = excel_writer.build_workbook(results)
    check("workbook produced", len(data) > 5000, len(data))
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data))
    check("workbook reopens", wb is not None)
    check("Summary sheet first", wb.sheetnames[0] == "Summary", wb.sheetnames)
    check("one Give and one Recv sheet per article",
          sum(1 for n in wb.sheetnames if n.startswith("Give")) == len(results)
          and sum(1 for n in wb.sheetnames if n.startswith("Recv")) == len(results),
          wb.sheetnames)
    check("all sheet names within Excel's 31-char limit",
          all(len(n) <= 31 for n in wb.sheetnames))
    formula_cells = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value[:1] in ("=", "+", "@"):
                    formula_cells.append((ws.title, cell.coordinate, cell.value[:20]))
    check("no unescaped formula cells", not formula_cells, formula_cells[:2])
    give = next(ws for ws in wb.worksheets if ws.title.startswith("Give"))
    headers = [c.value for c in give[1]]
    check("Give headers use the heuristic label",
          "Topic Overlap (heuristic)" in headers, headers)
    check("no column is labelled cannibalization",
          not any("cannibal" in (h or "").lower() for h in headers))


def test_input_guards():
    print("\nInput guards")
    for bad, why in [
        ("http://binaytara.org/x", "http rejected"),
        ("https://conference.binaytara.org/x", "subdomain rejected"),
        ("https://example.com/x", "external host rejected"),
        ("file:///etc/passwd", "file scheme rejected"),
        ("not a url", "garbage rejected"),
    ]:
        try:
            input_parser._assert_safe_url(bad)
            check(why, False, "was accepted")
        except input_parser.InputError:
            check(why, True)
    try:
        input_parser.from_docx(b"not a zip file at all", "x.docx")
        check("non-docx upload rejected", False)
    except input_parser.InputError:
        check("non-docx upload rejected", True)
    try:
        input_parser.from_docx(b"PK" + b"0" * (settings.MAX_UPLOAD_BYTES + 10), "x.docx")
        check("oversized upload rejected", False)
    except input_parser.InputError:
        check("oversized upload rejected", True)


def test_canonical_fix():
    """v3: verify the extractor returns _body_text for keyword scanner storage."""
    print("\nCanonical collision fix (v3)")
    rec = extract(ARTICLE_HTML, FIXTURE_URL)
    check("extractor returns _body_text", "_body_text" in rec, list(rec.keys()))
    check("_body_text is non-empty", len(rec.get("_body_text", "")) > 50,
          len(rec.get("_body_text", "")))
    check("canonical_url is captured",
          rec.get("canonical_url") == FIXTURE_URL, rec.get("canonical_url"))


if __name__ == "__main__":
    test_input_guards()
    test_canonical_fix()
    store = test_artifacts()
    if store is not None:
        html_res = test_html_pipeline(store)
        docx_res = test_docx_pipeline(store)
        test_workbook([html_res, docx_res])

    print("\n" + "=" * 60)
    if FAILURES:
        print(f"{len(FAILURES)} FAILURES: {FAILURES}")
        sys.exit(1)
    print("All end-to-end tests passed.")
