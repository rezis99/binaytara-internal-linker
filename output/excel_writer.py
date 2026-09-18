"""Formatted Excel export.

Every string written here comes from scraped pages or uploaded drafts, neither of
which is trusted input, so all cells are escaped against formula injection.
"""
from __future__ import annotations

import io
import re
from hashlib import blake2b

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BAND_FILL = {
    "High": PatternFill("solid", fgColor="E2EFDA"),
    "Medium": PatternFill("solid", fgColor="FFF2CC"),
    "Lower": PatternFill("solid", fgColor="F2F2F2"),
}
BAND_FONT = {"Lower": Font(color="808080")}
OVERLAP_FILL = {
    "High": PatternFill("solid", fgColor="FFC7CE"),
    "Medium": PatternFill("solid", fgColor="FFEB9C"),
}
OVERLAP_FONT = {
    "High": Font(color="9C0006", bold=True),
    "Medium": Font(color="9C5700"),
}
WRAP = Alignment(wrap_text=True, vertical="top")
LEFT_MARK = Border(left=Side(style="thick", color="1F4E79"))

# v4 (review item #9): decision columns first. A reviewer reads left to right
# and decides yes/no before reading the sentence, so Relevance, Match Type and
# the competition flag lead. "Use?" is an empty column the reviewer fills in.
GIVE_COLS = [
    ("Use?", 8), ("Relevance", 11), ("Match Type", 17),
    ("Keyword Competition", 30), ("Anchor Text", 22),
    ("Target Page Title", 34), ("Target Page Link", 46), ("Section", 12),
    ("Existing Sentence", 58), ("Modified Sentence", 58),
    ("Notes", 42),
]
RECEIVE_COLS = [
    ("Use?", 8), ("Relevance", 11), ("Match Type", 17),
    ("Keyword Competition", 30), ("Anchor Text", 22),
    ("Source Page Title", 34), ("Source Page", 46), ("Section", 12),
    ("Existing Sentence", 58), ("Modified Sentence", 58),
    ("Notes", 42),
]

_DANGEROUS = ("=", "+", "-", "@", "\t", "\r")


def esc(value):
    """Excel treats a leading = + - @ as a formula. openpyxl does not escape
    this, and the workbook will be emailed between team members."""
    if value is None:
        return ""
    if not isinstance(value, str):
        return value
    return "'" + value if value[:1] in _DANGEROUS else value


def sheet_name(prefix: str, url: str) -> str:
    """Stable across runs: a hash, not a positional counter, so re-running a
    failed batch puts each article on the same sheet as before."""
    slug = re.sub(r"[^A-Za-z0-9]+", "-", url.rstrip("/").split("/")[-1])[:18].strip("-")
    digest = blake2b(url.encode(), digest_size=2).hexdigest()
    name = f"{prefix} - {slug}-{digest}"
    return re.sub(r"[:\\/?*\[\]]", "-", name)[:31]


def _header(ws, cols) -> None:
    for i, (title, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=title)
        c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, WRAP
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"


def _style_row(ws, row_i: int, n_cols: int, band: str, overlap_col: int,
               overlap_level: str, recommended: bool, notes_col: int) -> None:
    fill = BAND_FILL.get(band)
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row_i, column=c)
        cell.alignment = WRAP
        if fill:
            cell.fill = fill
        if band in BAND_FONT:
            cell.font = BAND_FONT[band]
    if overlap_level in OVERLAP_FILL:
        oc = ws.cell(row=row_i, column=overlap_col)
        oc.fill = OVERLAP_FILL[overlap_level]
        oc.font = OVERLAP_FONT[overlap_level]
    if recommended:
        ws.cell(row=row_i, column=notes_col).border = LEFT_MARK
    ws.row_dimensions[row_i].height = 30


def _link(ws, row_i: int, col: int, url: str) -> None:
    cell = ws.cell(row=row_i, column=col, value=esc(url))
    cell.hyperlink = url
    cell.font = Font(color="0563C1", underline="single")
    cell.alignment = WRAP


def _competition_text(r) -> str:
    level = r.get("overlap_level") or "None"
    basis = r.get("overlap_basis") or "title heuristic"
    if level in ("", "None"):
        return f"None ({basis})"
    why = r.get("overlap_why") or ""
    return f"{level} ({basis}): {why}" if why else f"{level} ({basis})"


def write_give(ws, rows: list[dict]) -> None:
    _header(ws, GIVE_COLS)
    for i, r in enumerate(rows, start=2):
        values = [
            None,                                  # Use? left blank for the reviewer
            r["relevance"], r["match_type"], _competition_text(r), r["anchor"],
            r["target_title"], None, r["section"],
            r["existing_sentence"], r["modified_sentence"], r["notes"],
        ]
        for c, v in enumerate(values, 1):
            if c == 7:
                _link(ws, i, 7, r["target_url"])
            else:
                ws.cell(row=i, column=c, value=esc(v))
        ws.cell(row=i, column=5).font = Font(bold=True)          # anchor
        _style_row(ws, i, len(GIVE_COLS), r["relevance"], 4, r["overlap_level"],
                   r["notes"].startswith("Recommended"), 11)


def write_receive(ws, rows: list[dict], target_url: str) -> None:
    _header(ws, RECEIVE_COLS)
    ws.insert_rows(1)
    note = ws.cell(row=1, column=1,
                   value=esc(f"All rows below suggest linking TO: {target_url}"))
    note.font = Font(bold=True, italic=True)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(RECEIVE_COLS))}2"
    for i, r in enumerate(rows, start=3):
        values = [
            None,
            r["relevance"], r["match_type"], _competition_text(r), r["anchor"],
            r["source_title"], None, r["section"],
            r["existing_sentence"], r["modified_sentence"], r["notes"],
        ]
        for c, v in enumerate(values, 1):
            if c == 7:
                _link(ws, i, 7, r["source_url"])
            else:
                ws.cell(row=i, column=c, value=esc(v))
        ws.cell(row=i, column=5).font = Font(bold=True)
        _style_row(ws, i, len(RECEIVE_COLS), r["relevance"], 4, r["overlap_level"],
                   False, 11)


def write_summary(ws, results: list[dict]) -> None:
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 96
    row = 1
    for res in results:
        a, idx = res["article"], res["index"]
        give, recv = res["give"], res["receive"]
        pairs = [
            ("Article", a["title"]),
            ("Article URL" if not a["is_draft"] else "Uploaded draft", a["url"]),
            ("Word count", a["word_count"]),
            ("SOP benchmark link range", a["benchmark"]),
            ("Eligible paragraphs", a["eligible_paragraphs"]),
            ("Links to Give: total", len(give)),
            ("  of which High", sum(1 for r in give if r["relevance"] == "High")),
            ("  of which Medium", sum(1 for r in give if r["relevance"] == "Medium")),
            ("  of which Lower", sum(1 for r in give if r["relevance"] == "Lower")),
            ("Links to Receive: total", len(recv)),
            ("Existing internal links in body", len(a["existing_links"])),
            ("Date analysed", res["generated_at"]),
            ("Page database built", idx["built_at"][:10]),
            ("Pages indexed", idx["pages"]),
        ]
        for k, v in pairs:
            ws.cell(row=row, column=1, value=esc(k)).font = Font(bold=True)
            ws.cell(row=row, column=2, value=esc(v)).alignment = WRAP
            row += 1
        for link in a["existing_links"]:
            ws.cell(row=row, column=1, value="  existing link")
            ws.cell(row=row, column=2, value=esc(link))
            row += 1
        row += 2


def build_workbook(results: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    write_summary(ws, results)
    for res in results:
        url = res["article"]["url"]
        give = res["give"]
        topical = [r for r in give if r.get("is_topical", True)]
        other = [r for r in give if not r.get("is_topical", True)]
        write_give(wb.create_sheet(sheet_name("Give", url)), topical)
        if other:
            write_give(wb.create_sheet(sheet_name("Other", url)), other)
        write_receive(wb.create_sheet(sheet_name("Recv", url)), res["receive"], url)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
